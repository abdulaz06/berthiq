"""Bounded upload parsing, explicit mapping, canonical validation and persistence."""

from __future__ import annotations

import csv
import io
import json
import math
import uuid
import zipfile
from collections import Counter, defaultdict
from datetime import UTC, datetime
from pathlib import Path
from typing import Any, Literal
from xml.etree.ElementTree import ParseError

import sqlalchemy as sa
from defusedxml.common import DefusedXmlException
from lxml.etree import XMLSyntaxError
from sqlalchemy.orm import Session

from app.db.models import Case, CaseOutcome, Dataset, DatasetSource, DatasetStatus, Event
from app.schemas.datasets import (
    ColumnMapping,
    DataQualityReport,
    DatasetSummary,
    ValidationIssue,
)
from app.services.discovery import EventLite
from app.services.prediction import STAGES

MAX_BYTES = 5 * 1024 * 1024
MAX_ROWS = 50_000
MAX_COLUMNS = 64
PORT_PATH = Path(__file__).resolve().parents[2] / "data" / "ports" / "khalifa.json"


def port_config() -> dict[str, Any]:
    return json.loads(PORT_PATH.read_text(encoding="utf-8"))


def parse_upload(content: bytes, filename: str) -> list[dict[str, str]]:
    if len(content) > MAX_BYTES:
        raise ValueError("File exceeds the 5 MiB upload limit.")
    extension = Path(filename).suffix.lower()
    matrix: Any
    workbook = None
    if extension == ".csv":
        try:
            matrix = csv.reader(io.StringIO(content.decode("utf-8-sig")), strict=True)
        except UnicodeError as exc:
            raise ValueError("CSV must use UTF-8 encoding.") from exc
    elif extension == ".xlsx":
        from openpyxl import load_workbook

        try:
            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                if (
                    len(archive.infolist()) > 2000
                    or sum(item.file_size for item in archive.infolist()) > 30 * 1024 * 1024
                ):
                    raise ValueError("Workbook expands beyond the safe parsing limit.")
            workbook = load_workbook(
                io.BytesIO(content), read_only=True, data_only=True, keep_links=False
            )
            sheet = workbook.active
            if (
                sheet is None
                or (sheet.max_column or 0) > MAX_COLUMNS
                or (sheet.max_row or 0) > MAX_ROWS + 1
            ):
                workbook.close()
                raise ValueError("Workbook dimensions exceed the row/column limits.")
            matrix = sheet.iter_rows(values_only=True)
        except (
            zipfile.BadZipFile,
            KeyError,
            OSError,
            ParseError,
            XMLSyntaxError,
            DefusedXmlException,
        ) as exc:
            if workbook:
                workbook.close()
            raise ValueError("The uploaded file is not a readable XLSX workbook.") from exc
    else:
        raise ValueError("Choose a CSV or XLSX file.")
    try:
        headers = [str(v).strip() if v is not None else "" for v in next(matrix, [])]
        if not headers or len(headers) > MAX_COLUMNS or any(not h for h in headers):
            raise ValueError("A header row with 1–64 nonempty column names is required.")
        if len(headers) != len(set(headers)):
            raise ValueError("Column names must be unique.")
        rows: list[dict[str, str]] = []
        for values in matrix:
            if len(rows) >= MAX_ROWS:
                raise ValueError("The upload exceeds 50,000 event rows.")
            if len(values) > len(headers):
                raise ValueError("A row has more cells than the header.")
            row = {
                h: str(v).strip() if v is not None else ""
                for h, v in zip(headers, values, strict=False)
            }
            if any(len(v) > 10_000 for v in row.values()):
                raise ValueError("A cell exceeds the 10,000 character limit.")
            if any("\x00" in v for v in row.values()):
                raise ValueError("Cells must not contain NUL characters.")
            if any(row.values()):
                rows.append({h: row.get(h, "") for h in headers})
        if not rows:
            raise ValueError("The file contains no event rows.")
        return rows
    except (csv.Error, ParseError, XMLSyntaxError, DefusedXmlException) as exc:
        raise ValueError("File parsing failed; check its format and cell contents.") from exc
    finally:
        if workbook:
            workbook.close()


def suggest_mapping(columns: list[str]) -> ColumnMapping:
    aliases = {
        "case_id": ("case_id", "case", "caseid", "vessel", "port_call_id"),
        "activity": ("activity", "stage", "event", "event_name"),
        "end_ts": ("end_ts", "timestamp", "completed_at", "end_time"),
        "start_ts": ("start_ts", "started_at", "start_time"),
        "resource_id": ("resource_id", "resource", "reviewer"),
        "status": ("status", "state"),
    }
    normalized = {c.lower().strip().replace(" ", "_"): c for c in columns}
    fields = {
        field: next((normalized[a] for a in names if a in normalized), None)
        for field, names in aliases.items()
    }
    for required in ("case_id", "activity", "end_ts"):
        fields[required] = fields[required] or ""
    attrs = {
        c: c
        for c in columns
        if c
        in {
            "vessel_type",
            "loa",
            "draft",
            "cargo_type",
            "berth",
            "eta",
            "ata",
            "service_window_start",
            "service_window_end",
            "missing_doc_count",
            "weather_severity",
            "attributes",
        }
    }
    return ColumnMapping(**fields, attribute_map=attrs)


def timestamp(value: Any) -> datetime | None:
    if value in (None, ""):
        return None
    result = datetime.fromisoformat(str(value).replace("Z", "+00:00"))
    return result.replace(tzinfo=UTC) if result.tzinfo is None else result


def safe_attributes(value: dict[str, Any]) -> None:
    """Bound nested JSON before it reaches PostgreSQL or evidence serialization."""
    pending: list[tuple[Any, int]] = [(value, 0)]
    while pending:
        item, depth = pending.pop()
        if depth > 16:
            raise ValueError("Attribute JSON exceeds the nesting limit.")
        if isinstance(item, dict):
            pending.extend((v, depth + 1) for v in item.values())
            pending.extend((k, depth + 1) for k in item)
        elif isinstance(item, list):
            pending.extend((v, depth + 1) for v in item)
        elif isinstance(item, float) and not math.isfinite(item):
            raise ValueError("Non-finite JSON number.")
        elif isinstance(item, str) and "\x00" in item:
            raise ValueError("NUL character in JSON.")


def normalize(
    rows: list[dict[str, str]],
    mapping: ColumnMapping,
    dataset_id: str,
) -> tuple[list[EventLite], DataQualityReport]:
    issues: dict[str, tuple[Literal["fatal", "warning", "info"], str, list[int]]] = {}

    def issue(code: str, severity: Literal["fatal", "warning", "info"], msg: str, row: int):
        if code not in issues:
            issues[code] = (severity, msg, [])
        issues[code][2].append(row)

    cfg = port_config()
    known_resources = {"system", *(b["id"] for b in cfg["berths"])}
    for group in ("reviewers", "planners", "pilots", "agents"):
        known_resources.update(cfg[group]["pool"])
    events = []
    seen: set[str] = set()
    last: dict[str, datetime] = {}
    for index, row in enumerate(rows, start=2):
        cid, activity = row.get(mapping.case_id, ""), row.get(mapping.activity, "")
        if not cid or not activity:
            issue("missing_required", "fatal", "Case ID and activity are required.", index)
            continue
        try:
            end = timestamp(row.get(mapping.end_ts))
            start = timestamp(row.get(mapping.start_ts)) if mapping.start_ts else None
            if end is None:
                raise ValueError()
        except ValueError:
            issue("invalid_timestamp", "fatal", "Use valid ISO-8601 completion timestamps.", index)
            continue
        raw_end = row.get(mapping.end_ts, "")
        if datetime.fromisoformat(raw_end.replace("Z", "+00:00")).tzinfo is None:
            issue(
                "assumed_utc",
                "warning",
                "Timestamps without offsets are interpreted as UTC.",
                index,
            )
        if start and start > end:
            issue("negative_duration", "warning", "Negative duration; start time discarded.", index)
            start = None
        if cid in last and end < last[cid]:
            issue("impossible_order", "warning", "Rows are reordered by completion time.", index)
        last[cid] = end
        resource = row.get(mapping.resource_id, "") if mapping.resource_id else ""
        if resource and resource not in known_resources:
            issue(
                "unknown_resource",
                "warning",
                "Resource is not in the Khalifa configuration.",
                index,
            )
        attrs: dict[str, Any] = {}
        try:
            for source, key in mapping.attribute_map.items():
                value = row.get(source, "")
                if not value:
                    continue
                if key == "attributes":
                    parsed = json.loads(value)
                    if not isinstance(parsed, dict):
                        raise ValueError()
                    safe_attributes(parsed)
                    attrs.update(parsed)
                else:
                    attrs[key] = value
            for key in ("loa", "draft", "weather_severity", "missing_doc_count"):
                if key in attrs:
                    attrs[key] = float(attrs[key])
                    if not (-1e9 < attrs[key] < 1e9):
                        raise ValueError()
            for key in ("eta", "ata", "service_window_start", "service_window_end"):
                if key in attrs:
                    timestamp(attrs[key])
        except (ValueError, TypeError, RecursionError):
            issue(
                "invalid_attribute", "fatal", "Invalid maritime number, timestamp or JSON.", index
            )
            continue
        status = row.get(mapping.status) if mapping.status else None
        signature = json.dumps(
            [cid, activity, end.isoformat(), resource, status, attrs], sort_keys=True
        )
        if signature in seen:
            issue("duplicate_event", "warning", "Exact duplicate event excluded.", index)
            continue
        seen.add(signature)
        events.append(EventLite(cid, activity, start, end, resource or None, status, index, attrs))
    counts = Counter(e.case_id for e in events)
    for e in events:
        if counts[e.case_id] == 1:
            issue("single_event_case", "warning", "Case has only one observed event.", e.seq or 0)
    events.sort(key=lambda e: (e.case_id, e.end_ts, e.seq or 0))
    ordinals: Counter[str] = Counter()
    for e in events:
        ordinals[e.case_id] += 1
        e.seq = ordinals[e.case_id]
    report = DataQualityReport(
        dataset_id=dataset_id,
        total_events=len(rows),
        total_cases=len(counts),
        is_blocking=any(s == "fatal" for s, _, _ in issues.values()) or not events,
        issues=[
            ValidationIssue(severity=s, code=c, message=m, count=len(ids), sample_rows=ids[:8])
            for c, (s, m, ids) in sorted(issues.items())
        ],
    )
    return events, report


def persist_events(session: Session, dataset: Dataset, events: list[EventLite]) -> None:
    groups: dict[str, list[EventLite]] = defaultdict(list)
    for e in events:
        groups[e.case_id].append(e)
        session.add(Event(dataset_id=dataset.id, **vars(e)))
    active = 0
    for cid, evs in groups.items():
        attrs = {k: v for e in evs for k, v in e.attributes.items()}
        terminal = next((e for e in reversed(evs) if e.activity == STAGES[-1]), None)
        end_window = timestamp(attrs.get("service_window_end"))
        outcome = (
            CaseOutcome.active
            if terminal is None
            else (
                CaseOutcome.late
                if end_window and terminal.end_ts > end_window
                else CaseOutcome.on_time
            )
        )
        active += terminal is None
        session.add(
            Case(
                dataset_id=dataset.id,
                case_id=cid,
                is_active=terminal is None,
                current_stage=evs[-1].activity,
                first_event_ts=evs[0].end_ts,
                last_event_ts=evs[-1].end_ts,
                service_window_start=timestamp(attrs.get("service_window_start")),
                service_window_end=end_window,
                actual_berthing_ts=terminal.end_ts if terminal else None,
                outcome=outcome,
                rework_count=sum(e.status in {"returned", "rejected"} for e in evs),
                vessel_type=attrs.get("vessel_type"),
                loa=attrs.get("loa"),
                draft=attrs.get("draft"),
                cargo_type=attrs.get("cargo_type"),
                attributes=attrs,
            )
        )
    dataset.event_count, dataset.case_count = len(events), len(groups)
    dataset.time_range_start = min(e.end_ts for e in events)
    dataset.time_range_end = max(e.end_ts for e in events)
    dataset.stats = {
        **(dataset.stats or {}),
        "active_case_count": active,
        "activities": sorted({e.activity for e in events}),
        "resources": sorted({e.resource_id for e in events if e.resource_id}),
    }
    dataset.status = DatasetStatus.ready


def summary(dataset: Dataset) -> DatasetSummary:
    stats = dataset.stats or {}
    return DatasetSummary(
        id=str(dataset.id),
        name=dataset.name,
        source=dataset.source.value,
        status=dataset.status.value,
        port_config=dataset.port_config,
        event_count=dataset.event_count,
        case_count=dataset.case_count,
        active_case_count=stats.get("active_case_count", 0),
        time_range=(dataset.time_range_start, dataset.time_range_end)
        if dataset.time_range_start and dataset.time_range_end
        else None,
        activities=stats.get("activities", []),
        resources=stats.get("resources", []),
    )


def create_upload(
    session: Session,
    rows: list[dict[str, str]],
    name: str,
    owner: uuid.UUID | None = None,
) -> Dataset:
    dataset = Dataset(
        id=uuid.uuid4(),
        name=name[:160],
        source=DatasetSource.uploaded,
        status=DatasetStatus.validating,
        port_config="khalifa",
        event_count=0,
        case_count=0,
        stats={"raw_rows": rows},
        owner_token=owner,
    )
    session.add(dataset)
    session.commit()
    return dataset


def apply_mapping(session: Session, dataset: Dataset, mapping: ColumnMapping) -> DataQualityReport:
    if dataset.source != DatasetSource.uploaded or dataset.status == DatasetStatus.ready:
        raise ValueError("Ready datasets are immutable. Upload a new version to change mapping.")
    rows = (dataset.stats or {}).get("raw_rows", [])
    events, report = normalize(rows, mapping, str(dataset.id))
    dataset.stats = {
        **(dataset.stats or {}),
        "mapping": mapping.model_dump(),
        "validation": report.model_dump(mode="json"),
    }
    dataset.status = DatasetStatus.invalid if report.is_blocking else DatasetStatus.ready
    if not report.is_blocking:
        persist_events(session, dataset, events)
        dataset.stats = {k: v for k, v in dataset.stats.items() if k != "raw_rows"}
    session.commit()
    return report


def validate_dataset(session: Session, dataset: Dataset) -> DataQualityReport:
    stats = dataset.stats or {}
    if stats.get("validation"):
        return DataQualityReport.model_validate(stats["validation"])
    events = session.scalars(
        sa.select(Event).where(Event.dataset_id == dataset.id).order_by(Event.case_id, Event.seq)
    ).all()
    rows = [
        {
            "case_id": e.case_id,
            "activity": e.activity,
            "end_ts": e.end_ts.isoformat(),
            "start_ts": e.start_ts.isoformat() if e.start_ts else "",
            "resource_id": e.resource_id or "",
            "status": e.status or "",
            "attributes": json.dumps(e.attributes),
        }
        for e in events
    ]
    _, report = normalize(rows, suggest_mapping(list(rows[0]) if rows else []), str(dataset.id))
    return report
